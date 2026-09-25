import os
import re
import shutil
import datetime
import traceback
from flask import Flask, request, render_template, send_file, redirect, url_for, flash, session

import openpyxl
from openpyxl.cell.cell import MergedCell

import transform_type1 as t1
import transform_type2 as t2
import auth

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(BASE_DIR)
PNL_PATH = os.path.join(PROJECT_DIR, "P&L STATEMENT.xlsx")
ACPLACT_PATH = os.path.join(PROJECT_DIR, "acplact.xlsx")

UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
OUTPUT_DIR = os.path.join(BASE_DIR, "outputs")
HISTORY_DIR = os.path.join(BASE_DIR, "history")
PENDING_DIR = os.path.join(BASE_DIR, "pending")
PENDING_PNL_PATH = os.path.join(PENDING_DIR, "pending_pnl.xlsx")
PENDING_ACPLACT_PATH = os.path.join(PENDING_DIR, "pending_acplact.xlsx")

for d in (UPLOAD_DIR, OUTPUT_DIR, HISTORY_DIR, PENDING_DIR):
    os.makedirs(d, exist_ok=True)

app = Flask(__name__)
app.secret_key = "astra-chemtech-pnl-tool"  # signs the login session cookie; kept stable across restarts on purpose

PUBLIC_ROUTES = {"login"}  # accessible without being logged in


@app.before_request
def require_login():
    if request.endpoint is None:
        return  # e.g. a 404 on a non-matching path — let it 404 normally
    if request.endpoint in PUBLIC_ROUTES or request.endpoint == "static":
        return
    if not session.get("logged_in"):
        return redirect(url_for("login", next=request.path))


def _safe_next_path(candidate):
    """Only ever redirect to a path within this app — a `next` value from
    a query string or form field is user-controlled, so an unrestricted
    redirect there would be an open-redirect vector."""
    if candidate and candidate.startswith("/") and not candidate.startswith("//"):
        return candidate
    return url_for("index")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        if auth.verify_login(username, password):
            session["logged_in"] = True
            session["username"] = username
            return redirect(_safe_next_path(request.form.get("next")))
        flash("Incorrect username or password.", "error")
    next_path = request.args.get("next", "")
    return render_template("login.html", next_path=next_path)


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/change-password", methods=["GET", "POST"])
def change_password():
    if request.method == "POST":
        old_password = request.form.get("old_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")
        if new_password != confirm_password:
            flash("New password and confirmation don't match.", "error")
        else:
            ok, error = auth.change_password(old_password, new_password)
            if ok:
                flash("Password changed.", "ok")
                return redirect(url_for("index"))
            flash(error, "error")
    return render_template("change_password.html")

TRANSFORMATIONS = {
    "pnl": {
        "label": "PNL",
        "description": "The monthly P&L STATEMENT.xlsx merge — matches Sheet1 by GL Code, adds the new month as a column.",
    },
    "acplact": {
        "label": "ACPLACT",
        "description": "The FIFO-sch / FIFO-PNL merge in acplact.xlsx — GL Code (incl. merged multi-code rows) summed from the same Sheet1.",
    },
}

MODES = {
    "new_month": {
        "label": "New month",
        "description": "Add this month as a brand-new column, shifting older months right.",
    },
    "replace_latest": {
        "label": "Replace latest month",
        "description": "Overwrite the current latest month's data with corrected/final figures — use this when more data arrived after you already processed that month.",
    },
}

SOURCE_LABELS = {
    "matched": "From Sheet1 (matched by GL Code)",
    "zeroed": "Set to 0 — not found in this Sheet1",
    "new_gl": "NEW row inserted — check it belongs here or add its code to an existing row",
    "template": "Auto-filled — subtotal/movement formula",
    "fiscal-accrual": "Auto-filled — fiscal-YTD formula, multiplier bumped",
    "flat-value": "Auto-filled — unchanged from prior months",
    "left-blank": "Left blank — needs manual entry",
    "manual-formula-left-blank": "Left blank — old formula had no safe pattern",
    "no-formula-found": "Left blank — no formula to carry forward",
    "adjustment_dropped": "One-off adjustment dropped — review if one applies this month",
    "gross-sales-matched": "From Sheet1's Gross Sales row",
    "adopted": "GL Code attached to your existing row (no duplicate added)",
    "partial": "Some of this row's GL codes are missing from Sheet1",
}


def cell_display_value(v):
    if v is None:
        return ""
    return str(v)


def parse_edited_value(text):
    """Turn a review-page text field back into a cell value: blank -> None,
    '=...' -> kept as a formula string, else numeric if possible, else raw text."""
    text = text.strip()
    if text == "":
        return None
    if text.startswith("="):
        return text
    try:
        f = float(text)
        return int(f) if f.is_integer() else f
    except ValueError:
        return text


def copy_sheet_values(src_ws, dest_ws):
    """Copy values (not formulas' side-effects) from an uploaded sheet into
    the persisted workbook's Sheet1, skipping merged-cell followers."""
    for row in src_ws.iter_rows():
        for cell in row:
            target = dest_ws.cell(row=cell.row, column=cell.column)
            if isinstance(target, MergedCell):
                continue
            target.value = cell.value


NEEDS_ATTENTION_TAGS = {
    "zeroed", "left-blank", "manual-formula-left-blank", "no-formula-found", "adjustment_dropped",
    "new_gl", "partial",
}


def build_pnl_review_rows(ws, report):
    """One row per GL-linked or manual/subtotal row touched by the merge,
    editable value/formula for the new/latest month's column L."""
    t1.detect_layout(ws)
    new_codes = {code for code, _ in report.get("new_gl_rows_inserted", [])}
    adopted_codes = {code for code, _, _ in report.get("adopted_rows", [])}
    partial_rows = {r for r, _, _ in report.get("gl_rows_partial", [])}
    manual_tags = {r: tag for r, _, tag in report.get("manual_rows", [])}
    suggestions = report.get("suggestions", {})

    rows = []
    for r in range(1, ws.max_row + 1):
        code = ws.cell(row=r, column=t1.GL_COL).value
        desc = ws.cell(row=r, column=t1.DESC_COL).value
        codes = t1.parse_codes(code)
        if r in manual_tags:
            source = manual_tags[r]
        elif code and r > t1.HEADER_ROW:
            if any(c in new_codes for c in codes):
                source = "new_gl"
            elif any(c in adopted_codes for c in codes):
                source = "adopted"
            elif r in partial_rows:
                source = "partial"
            else:
                source = "matched"
        else:
            continue

        cell = ws.cell(row=r, column=t1.TOTAL_COL)
        value = cell_display_value(cell.value)
        if value == "" and r in suggestions:
            value = suggestions[r]
        rows.append({
            "row": r,
            "code": code or "",
            "desc": desc or "",
            "value": value,
            "source": SOURCE_LABELS.get(source, source),
            "needs_attention": source in NEEDS_ATTENTION_TAGS,
        })
    return rows


def build_acplact_editable_rows(wb, report):
    """Only the fields ACPLACT can never get from Sheet1 at all — GL codes
    with no match anywhere in the trial balance, non-GL rows left blank
    with no safe pattern, and Gross Sales (which the client types directly
    into Sheet1's label row). Calculated rows (subtotals, GST, the
    Opening/Closing-stock rollforward) are never made editable here."""
    sch_ws = wb[t2.FIFO_SCH_SHEET]
    pnl_ws = wb[t2.FIFO_PNL_SHEET]
    sch_amount_col = t2._current_amount_col(sch_ws, t2.SCH_HEADER_ROW, t2.SCH_ANCHOR_COL)
    pnl_amount_col = t2._current_amount_col(pnl_ws, t2.PNL_HEADER_ROW, t2.PNL_ANCHOR_COL)

    manual_row_nums = {r for r, _code, _desc in report.get("gl_rows_manual", [])}
    left_blank_rows = {r for r, _desc, tag in report.get("manual_rows", []) if tag == "left-blank"}
    editable_sch_rows = manual_row_nums | left_blank_rows
    suggestions = report.get("suggestions", {})

    rows = []
    for r in sorted(editable_sch_rows):
        cell = sch_ws.cell(row=r, column=sch_amount_col)
        # the cell itself is deliberately left blank (we don't trust the old
        # formula enough to auto-apply it) — but show it as a starting point
        # to edit rather than an empty box with no context
        display_value = cell_display_value(cell.value)
        if display_value == "" and r in suggestions:
            display_value = str(suggestions[r])
        rows.append({
            "sheet": "sch",
            "row": r,
            "code": sch_ws.cell(row=r, column=t2.SCH_GL_COL).value or "",
            "desc": sch_ws.cell(row=r, column=t2.SCH_DESC_COL).value or "",
            "value": display_value,
            "needs_attention": True,
        })

    gross_cell = pnl_ws.cell(row=t2.PNL_GROSS_SALES_ROW, column=pnl_amount_col)
    rows.append({
        "sheet": "pnl",
        "row": t2.PNL_GROSS_SALES_ROW,
        "code": "",
        "desc": pnl_ws.cell(row=t2.PNL_GROSS_SALES_ROW, column=t2.PNL_DESC_COL).value or "Gross Sales",
        "value": cell_display_value(gross_cell.value),
        "needs_attention": True,
    })
    return rows


def get_pnl_current_month():
    if not os.path.isfile(PNL_PATH):
        return None
    try:
        wb = openpyxl.load_workbook(PNL_PATH, data_only=True)
        d = t1.get_current_pnl_month(wb["P&L"])
        return d.strftime("%B %Y") if d else None
    except Exception:
        return None


def get_acplact_current_month():
    if not os.path.isfile(ACPLACT_PATH):
        return None
    try:
        wb = openpyxl.load_workbook(ACPLACT_PATH, data_only=True)
        d = t2.get_current_month(wb)
        return d.strftime("%B %Y") if d else None
    except Exception:
        return None


BACKUP_TS_RE = re.compile(r'_(\d{8}_\d{6})\.xlsx$')


def latest_backup_batch():
    """PNL and ACPLACT backups made in the same /confirm call share the
    same timestamp suffix (e.g. '20260917_193900') — group by that so
    Undo can restore every file from the most recent run together,
    instead of only whichever single file happens to be newest by mtime."""
    names = [f for f in os.listdir(HISTORY_DIR) if f.endswith(".xlsx")]
    with_ts = [(BACKUP_TS_RE.search(f).group(1), f) for f in names if BACKUP_TS_RE.search(f)]
    if not with_ts:
        return []
    latest_ts = max(ts for ts, _ in with_ts)
    return [os.path.join(HISTORY_DIR, f) for ts, f in with_ts if ts == latest_ts]


def recent_outputs():
    """Newest finished file of each kind, so the client can re-download last
    time's result on a fresh visit."""
    found = []
    for label, prefix in (("PNL", "P&L STATEMENT ("), ("ACPLACT", "acplact (")):
        files = [f for f in os.listdir(OUTPUT_DIR) if f.startswith(prefix) and f.endswith(".xlsx")]
        if not files:
            continue
        newest = max(files, key=lambda f: os.path.getmtime(os.path.join(OUTPUT_DIR, f)))
        when = datetime.datetime.fromtimestamp(os.path.getmtime(os.path.join(OUTPUT_DIR, newest)))
        found.append({"label": label, "filename": newest, "when": when.strftime("%d %b %Y, %H:%M")})
    return found


@app.route("/", methods=["GET"])
def index():
    pnl_month = get_pnl_current_month()
    acplact_month = get_acplact_current_month()
    can_undo = len(latest_backup_batch()) > 0
    return render_template(
        "index.html",
        transformations=TRANSFORMATIONS,
        modes=MODES,
        pnl_month=pnl_month,
        acplact_month=acplact_month,
        can_undo=can_undo,
        recent=recent_outputs(),
    )


@app.route("/process", methods=["POST"])
def process():
    file = request.files.get("sheet1_file")
    if not file or file.filename == "":
        flash("Please choose an Excel file to upload.", "error")
        return redirect(url_for("index"))

    selected = request.form.getlist("transforms")
    if not selected:
        flash("Please select at least one transformation to run.", "error")
        return redirect(url_for("index"))

    mode_key = request.form.get("mode", "new_month")
    if mode_key not in MODES:
        flash("Unknown mode selected.", "error")
        return redirect(url_for("index"))

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    upload_path = os.path.join(UPLOAD_DIR, f"sheet1_{ts}.xlsx")
    file.save(upload_path)

    sections = []

    try:
        uploaded_wb = openpyxl.load_workbook(upload_path, data_only=True)
        uploaded_ws = uploaded_wb["Sheet1"] if "Sheet1" in uploaded_wb.sheetnames \
            else uploaded_wb.worksheets[0]

        if "pnl" in selected:
            if not os.path.isfile(PNL_PATH):
                raise FileNotFoundError("No existing P&L STATEMENT.xlsx found in the project folder.")
            pnl_wb = openpyxl.load_workbook(PNL_PATH)  # keep formulas
            if "Sheet1" not in pnl_wb.sheetnames:
                pnl_wb.create_sheet("Sheet1")
            copy_sheet_values(uploaded_ws, pnl_wb["Sheet1"])
            run_fn = t1.apply_new_month if mode_key == "new_month" else t1.replace_latest_month
            report = run_fn(pnl_wb)
            pnl_wb.save(PENDING_PNL_PATH)
            sections.append({
                "key": "pnl",
                "label": TRANSFORMATIONS["pnl"]["label"],
                "report": report,
                "rows": build_pnl_review_rows(pnl_wb["P&L"], report),
                "editable": True,
            })

        if "acplact" in selected:
            if not os.path.isfile(ACPLACT_PATH):
                raise FileNotFoundError("No existing acplact.xlsx found in the project folder.")
            acplact_wb = openpyxl.load_workbook(ACPLACT_PATH)
            run_fn = t2.apply_new_month if mode_key == "new_month" else t2.replace_latest_month
            report = run_fn(acplact_wb, uploaded_ws)
            acplact_wb.save(PENDING_ACPLACT_PATH)
            sections.append({
                "key": "acplact",
                "label": TRANSFORMATIONS["acplact"]["label"],
                "report": report,
                "rows": [],
                "editable": False,
                "editable_rows": build_acplact_editable_rows(acplact_wb, report),
            })

        return render_template("review.html", sections=sections, mode_key=mode_key,
                                mode_label=MODES[mode_key]["label"],
                                source_labels=SOURCE_LABELS)

    except (t1.NotANewMonthError, t1.NotTheLatestMonthError,
            t2.NotANewMonthError, t2.NotTheLatestMonthError) as e:
        for p in (PENDING_PNL_PATH, PENDING_ACPLACT_PATH):
            if os.path.isfile(p):
                os.remove(p)
        flash(str(e), "error")
        return redirect(url_for("index"))

    except Exception as e:
        for p in (PENDING_PNL_PATH, PENDING_ACPLACT_PATH):
            if os.path.isfile(p):
                os.remove(p)
        flash(f"Processing failed: {e}", "error")
        traceback.print_exc()
        return redirect(url_for("index"))


@app.route("/confirm", methods=["POST"])
def confirm():
    mode_key = request.form.get("mode_key", "new_month")
    ran_pnl = os.path.isfile(PENDING_PNL_PATH)
    ran_acplact = os.path.isfile(PENDING_ACPLACT_PATH)

    if not ran_pnl and not ran_acplact:
        flash("Nothing pending to confirm — please upload again.", "error")
        return redirect(url_for("index"))

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    downloads = []

    if ran_pnl:
        pnl_wb = openpyxl.load_workbook(PENDING_PNL_PATH)
        ws = pnl_wb["P&L"]
        t1.detect_layout(ws)
        edited = 0
        for key, text in request.form.items():
            if not key.startswith("cell_pnl_"):
                continue
            row_no = int(key[len("cell_pnl_"):])
            new_value = parse_edited_value(text)
            cell = ws.cell(row=row_no, column=t1.TOTAL_COL)
            if cell_display_value(cell.value) != text.strip():
                cell.value = new_value
                edited += 1

        new_month = ws.cell(row=7, column=t1.TOTAL_COL).value.date().isoformat()
        backup_path = os.path.join(HISTORY_DIR, f"P&L_before_{mode_key}_{ts}.xlsx")
        shutil.copy(PNL_PATH, backup_path)
        ws.column_dimensions["B"].hidden = False  # GL Code always visible
        pnl_wb.save(PNL_PATH)

        if "Sheet1" in pnl_wb.sheetnames:
            pnl_wb.remove(pnl_wb["Sheet1"])
        output_name = f"P&L STATEMENT ({new_month}).xlsx"
        output_path = os.path.join(OUTPUT_DIR, output_name)
        pnl_wb.save(output_path)
        downloads.append({"label": "PNL", "filename": output_name, "edited": edited})
        os.remove(PENDING_PNL_PATH)

    if ran_acplact:
        acplact_wb = openpyxl.load_workbook(PENDING_ACPLACT_PATH)
        sch_ws = acplact_wb[t2.FIFO_SCH_SHEET]
        pnl_ws = acplact_wb[t2.FIFO_PNL_SHEET]
        sch_amount_col = t2._current_amount_col(sch_ws, t2.SCH_HEADER_ROW, t2.SCH_ANCHOR_COL)
        pnl_amount_col = t2._current_amount_col(pnl_ws, t2.PNL_HEADER_ROW, t2.PNL_ANCHOR_COL)

        acplact_edited = 0
        for key, text in request.form.items():
            if key.startswith("cell_acplact_sch_"):
                ws, col = sch_ws, sch_amount_col
                row_no = int(key[len("cell_acplact_sch_"):])
            elif key.startswith("cell_acplact_pnl_"):
                ws, col = pnl_ws, pnl_amount_col
                row_no = int(key[len("cell_acplact_pnl_"):])
            else:
                continue
            new_value = parse_edited_value(text)
            cell = ws.cell(row=row_no, column=col)
            if cell_display_value(cell.value) != text.strip():
                cell.value = new_value
                acplact_edited += 1

        new_month = t2.get_current_month(acplact_wb).isoformat()
        backup_path = os.path.join(HISTORY_DIR, f"acplact_before_{mode_key}_{ts}.xlsx")
        shutil.copy(ACPLACT_PATH, backup_path)
        acplact_wb.save(ACPLACT_PATH)

        output_name = f"acplact ({new_month}).xlsx"
        output_path = os.path.join(OUTPUT_DIR, output_name)
        acplact_wb.save(output_path)
        downloads.append({"label": "ACPLACT", "filename": output_name, "edited": acplact_edited})
        os.remove(PENDING_ACPLACT_PATH)

    return render_template("result.html", downloads=downloads, mode_label=MODES[mode_key]["label"])


@app.route("/cancel", methods=["POST"])
def cancel():
    for p in (PENDING_PNL_PATH, PENDING_ACPLACT_PATH):
        if os.path.isfile(p):
            os.remove(p)
    flash("Discarded — nothing was saved.", "ok")
    return redirect(url_for("index"))


@app.route("/undo", methods=["POST"])
def undo():
    backups = latest_backup_batch()
    if not backups:
        flash("No previous version to restore to.", "error")
        return redirect(url_for("index"))

    reverted = []
    for backup in backups:
        name = os.path.basename(backup)
        if name.startswith("acplact_before_"):
            target, label = ACPLACT_PATH, "ACPLACT"
        else:
            target, label = PNL_PATH, "PNL"
        shutil.copy(backup, target)
        os.remove(backup)
        reverted.append(label)

    flash(f"Reverted {' and '.join(reverted)} to their previous version.", "ok")
    return redirect(url_for("index"))


@app.route("/download/<path:filename>")
def download(filename):
    filename = os.path.basename(filename)
    path = os.path.join(OUTPUT_DIR, filename)
    if not os.path.isfile(path):
        flash("File not found.", "error")
        return redirect(url_for("index"))
    return send_file(path, as_attachment=True, download_name=filename)


if __name__ == "__main__":
    app.run(debug=False, port=5050)
