"""Permanent record of every save (new month, replace, restore, undo), a
saved copy of each resulting workbook, and a check that flags any change to
figures of EARLIER months — so a wrong run can be spotted and a correct
earlier version brought back."""
import os
import re
import json
import shutil
import datetime
import openpyxl

import transform_type1 as t1
import transform_type2 as t2

MAX_CHANGES_KEPT = 300
TOL = 0.005


def _is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def _differs(old, new):
    """Only compares what can be compared reliably: two constants, or a
    constant vs. a formula. Two formulas are skipped because inserting rows
    or columns legitimately rewrites their cell references."""
    if _is_formula(old) and _is_formula(new):
        return False
    if _is_formula(old) != _is_formula(new):
        return True
    if isinstance(old, (int, float)) and isinstance(new, (int, float)):
        return abs(old - new) > TOL
    if old in (None, "") and new in (None, "", 0):
        return False
    if new in (None, "") and old in (None, "", 0):
        return False
    return old != new


def _row_keys(ws, code_col, desc_col, first_row):
    """row number -> stable key: (code or description, n-th time seen)."""
    seen, keys = {}, {}
    for r in range(first_row, ws.max_row + 1):
        code = ws.cell(row=r, column=code_col).value
        desc = ws.cell(row=r, column=desc_col).value
        base = str(code).strip() if code else (str(desc).strip() if desc else "")
        if not base:
            continue
        seen[base] = seen.get(base, 0) + 1
        keys[r] = (base, seen[base])
    return keys


def _pnl_cols(wb):
    ws = wb["P&L"]
    t1.detect_layout(ws)
    out = {}
    for c in t1.get_month_data_columns(ws):
        d = ws.cell(row=t1.DATE_ROW, column=c).value
        if isinstance(d, datetime.datetime):
            out[("", d.strftime("%Y-%m"))] = c
    return ws, out


def _sch_cols(ws):
    out, c = {}, t2.SCH_ANCHOR_COL
    while ws.cell(row=t2.SCH_DATE_ROW, column=c).value not in (None, ""):
        try:
            d = t2.parse_sch_date(ws.cell(row=t2.SCH_DATE_ROW, column=c).value)
        except ValueError:
            break
        kind = "movement" if str(ws.cell(row=t2.SCH_HEADER_ROW, column=c).value or "").strip().lower().startswith("only") else "amount"
        out[(kind, d.strftime("%Y-%m"))] = c
        c += 1
    return out


def _fpnl_cols(ws):
    out, c = {}, t2.PNL_ANCHOR_COL
    while ws.cell(row=t2.PNL_HEADER_ROW, column=c).value not in (None, ""):
        label = " ".join(str(ws.cell(row=t2.PNL_HEADER_ROW, column=c).value).split())
        kind = "movement" if label.lower().startswith("only") else "amount"
        m = re.search(r"([A-Za-z]+)-(\d{4})", label)
        if m:
            try:
                d = datetime.datetime.strptime(f"{m.group(1)} {m.group(2)}", "%B %Y")
                out[(kind, d.strftime("%Y-%m"))] = c
            except ValueError:
                pass
        c += 1
    return out


def _compare_sheet(name, ws_old, ws_new, cols_old, cols_new, keys_old, keys_new, skip_month, out):
    key_to_new_row = {k: r for r, k in keys_new.items()}
    for (kind, month), c_old in cols_old.items():
        if month == skip_month or (kind, month) not in cols_new:
            continue
        c_new = cols_new[(kind, month)]
        for r_old, key in keys_old.items():
            r_new = key_to_new_row.get(key)
            if r_new is None:
                continue
            old = ws_old.cell(row=r_old, column=c_old).value
            new = ws_new.cell(row=r_new, column=c_new).value
            if _differs(old, new):
                out.append({"sheet": name, "row": key[0], "month": month, "kind": kind,
                            "old": "" if old is None else str(old),
                            "new": "" if new is None else str(new)})


def earlier_month_changes(kind, before_path, after_path, mode_key):
    """Cells of months that existed before this run whose figures now differ.
    In 'replace' mode the newest month is being corrected on purpose, so it
    is left out of the comparison."""
    before = openpyxl.load_workbook(before_path)
    after = openpyxl.load_workbook(after_path)
    out = []
    if kind == "PNL":
        wo, co = _pnl_cols(before)
        ko = _row_keys(wo, t1.GL_COL, t1.DESC_COL, t1.FIRST_DATA_ROW)
        wn, cn = _pnl_cols(after)
        kn = _row_keys(wn, t1.GL_COL, t1.DESC_COL, t1.FIRST_DATA_ROW)
        skip = max((m for _, m in co), default=None) if mode_key.startswith("replace") else None
        _compare_sheet("P&L", wo, wn, co, cn, ko, kn, skip, out)
    else:
        so, sn = before[t2.FIFO_SCH_SHEET], after[t2.FIFO_SCH_SHEET]
        co, cn = _sch_cols(so), _sch_cols(sn)
        skip = max((m for _, m in co), default=None) if mode_key.startswith("replace") else None
        ko = _row_keys(so, t2.SCH_GL_COL, t2.SCH_DESC_COL, t2.SCH_FIRST_DATA_ROW)
        kn = _row_keys(sn, t2.SCH_GL_COL, t2.SCH_DESC_COL, t2.SCH_FIRST_DATA_ROW)
        _compare_sheet("FIFO-sch", so, sn, co, cn, ko, kn, skip, out)
        po, pn = before[t2.FIFO_PNL_SHEET], after[t2.FIFO_PNL_SHEET]
        co, cn = _fpnl_cols(po), _fpnl_cols(pn)
        skip = max((m for _, m in co), default=None) if mode_key.startswith("replace") else None
        ko = _row_keys(po, t2.PNL_GL_COL, t2.PNL_DESC_COL, t2.PNL_FIRST_DATA_ROW)
        kn = _row_keys(pn, t2.PNL_GL_COL, t2.PNL_DESC_COL, t2.PNL_FIRST_DATA_ROW)
        _compare_sheet("FIFO - P & L", po, pn, co, cn, ko, kn, skip, out)
    return out


class HistoryLog:
    def __init__(self, history_dir):
        self.dir = history_dir
        self.versions = os.path.join(history_dir, "versions")
        self.path = os.path.join(history_dir, "log.json")
        os.makedirs(self.versions, exist_ok=True)

    def entries(self):
        try:
            with open(self.path, encoding="utf-8") as f:
                return json.load(f)
        except (OSError, ValueError):
            return []

    def _save(self, entries):
        tmp = self.path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(entries, f, indent=1)
        os.replace(tmp, self.path)

    def get(self, entry_id):
        return next((e for e in self.entries() if e["id"] == entry_id), None)

    def add(self, kind, action, month, saved_path, ts, edited=0, source_file="",
            changes=None, note=""):
        """Copies the saved workbook into versions/ and appends a log line."""
        entries = self.entries()
        entry_id = f"{ts}_{kind}"
        ext_name = f"{entry_id}.xlsx"
        shutil.copy(saved_path, os.path.join(self.versions, ext_name))
        changes = changes or []
        entries.append({
            "id": entry_id, "kind": kind, "action": action, "month": month,
            "when": datetime.datetime.now().strftime("%d %b %Y, %H:%M:%S"),
            "edited": edited, "source_file": source_file, "note": note,
            "version_file": ext_name,
            "changed_count": len(changes), "changes": changes[:MAX_CHANGES_KEPT],
        })
        self._save(entries)
        return entry_id

    def version_path(self, entry):
        return os.path.join(self.versions, entry["version_file"])
